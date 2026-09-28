
from io import BytesIO

from docx import Document

from backend.formatters.documents import (
    format_txt,
    format_docx,
    format_pdf,
)


def test_txt():
    data = format_txt("Hello LegalEase")

    assert data.decode("utf-8") == "Hello LegalEase"


def test_docx():
    data = format_docx(
        "1. Confidentiality\nSign here",
        "NDA",
        "Term one;Term two",
    )

    doc = Document(BytesIO(data))

    assert any(
        "Confidentiality" in paragraph.text
        for paragraph in doc.paragraphs
    )

    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 3


def test_pdf():
    data = format_pdf(
        "1. Confidentiality\nSign here",
        "NDA",
        "Term one;Term two",
    )

    assert data.startswith(b"%PDF")