
"""Export editable legal drafts to TXT, DOCX and PDF."""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF


FONT_PATHS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
]


def get_terms(terms):
    return [
        term.strip()
        for term in terms.split(";")
        if term.strip()
    ]


def format_txt(text):
    return text.encode("utf-8")


def format_docx(text, doc_type, terms="", logo=None):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    if logo:
        doc.add_picture(
            BytesIO(logo),
            width=Inches(1.1),
        )

    doc.add_heading(doc_type, 0)

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        if line.startswith(("- ", "* ", "• ")):
            doc.add_paragraph(
                line[2:],
                style="List Bullet",
            )

        elif line.isupper() and len(line) < 95:
            doc.add_heading(line, level=2)

        else:
            doc.add_paragraph(line)

    items = get_terms(terms)

    if items:
        doc.add_heading(
            "User-specified terms (reference)",
            level=2,
        )

        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"

        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, item in enumerate(items, 1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = item

    footer = section.footer.paragraphs[0]
    footer.text = (
        "LegalEase | AI-generated draft - "
        "Legal review required"
    )

    buffer = BytesIO()
    doc.save(buffer)

    return buffer.getvalue()


class BrandedPDF(FPDF):
    def __init__(self, title, logo=None):
        super().__init__()

        self.document_title = title
        self.logo_bytes = logo

        self.set_auto_page_break(
            auto=True,
            margin=20,
        )

        font_path = next(
            (
                path
                for path in FONT_PATHS
                if Path(path).exists()
            ),
            None,
        )

        if font_path:
            self.add_font(
                "LegalUnicode",
                "",
                font_path,
            )
            self.body_font = "LegalUnicode"
        else:
            self.body_font = "Helvetica"

    def safe_text(self, value):
        if self.body_font == "Helvetica":
            return value.encode(
                "latin-1",
                "replace",
            ).decode("latin-1")

        return value

    def header(self):
        if self.logo_bytes:
            from tempfile import NamedTemporaryFile
            import os

            with NamedTemporaryFile(
                suffix=".png",
                delete=False,
            ) as temp:
                temp.write(self.logo_bytes)
                logo_path = temp.name

            try:
                self.image(
                    logo_path,
                    x=12,
                    y=9,
                    h=13,
                )
            finally:
                os.unlink(logo_path)

        self.set_font(
            self.body_font,
            size=10,
        )

        self.set_y(12)

        self.cell(
            0,
            8,
            self.safe_text(self.document_title),
            align="R",
        )

        self.ln(12)

    def footer(self):
        self.set_y(-15)

        self.set_font(
            self.body_font,
            size=8,
        )

        self.cell(
            0,
            8,
            (
                "LegalEase | Draft for legal review | "
                f"Page {self.page_no()}"
            ),
            align="C",
        )



def format_pdf(text, doc_type, terms="", logo=None):
    pdf = BrandedPDF(doc_type, logo)

    pdf.set_margins(
        left=15,
        top=30,
        right=15,
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    # Document title
    pdf.set_x(pdf.l_margin)
    pdf.set_font(pdf.body_font, size=16)

    pdf.multi_cell(
        w=0,
        h=10,
        text=pdf.safe_text(doc_type),
    )

    pdf.ln(4)

    # Document content
    pdf.set_font(pdf.body_font, size=11)

    for line in text.splitlines():
        line = line.strip()

        if not line:
            pdf.ln(4)
            continue

        # Remove problematic control characters
        line = "".join(
            char for char in line
            if char == "\t" or ord(char) >= 32
        )

        if not line:
            continue

        pdf.set_x(pdf.l_margin)

        pdf.multi_cell(
            w=0,
            h=6,
            text=pdf.safe_text(line),
            wrapmode="CHAR",
        )

    # User-specified terms
    items = get_terms(terms)

    if items:
        pdf.ln(5)

        pdf.set_x(pdf.l_margin)
        pdf.set_font(pdf.body_font, size=13)

        pdf.multi_cell(
            w=0,
            h=8,
            text="User-specified terms (reference)",
        )

        pdf.set_font(pdf.body_font, size=10)

        for index, term in enumerate(items, 1):
            term = "".join(
                char for char in term
                if char == "\t" or ord(char) >= 32
            )

            if not term:
                continue

            pdf.set_x(pdf.l_margin)

            pdf.multi_cell(
                w=0,
                h=6,
                text=pdf.safe_text(
                    f"{index}. {term}"
                ),
                wrapmode="CHAR",
            )

    return bytes(pdf.output())